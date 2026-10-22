# -*- coding: utf-8 -*-
"""ppt_embed.py — 把检索录屏装配进 PPT：底部深蓝条（缩略图/或直接视频内嵌）+ 路径说明
用法:
  python ppt_embed.py --ppt <pptx路径> --thumbs <thumbs目录> --spec <spec.json>
spec.json 形如:
{
  "3": {"videos_dir": "D:/proj/05_输出文档/检索录屏",
        "items": [{"key":"CNKI_S001","label":"S001·528","video":"录屏_CNKI_S001_检索过程_20261021.mp4"}],
        "path_text": "05_输出文档\\检索录屏\\录屏_CNKI_S001…S006_…mp4（6 段）"},
  "4": {...}, "5": {...}
}
键 = 1-based 页码。铁律：坐标全程 EMU；OneDrive 环境走 TEMP 副本（本脚本内置）。
"""
import sys, os, json, shutil, argparse
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from PIL import Image

DARK = RGBColor(0x1F, 0x4E, 0x79)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xCF, 0xE2, 0xF3)

ap = argparse.ArgumentParser()
ap.add_argument('--ppt', required=True)
ap.add_argument('--thumbs', required=True)
ap.add_argument('--spec', required=True)
ap.add_argument('--embed-video', action='store_true', help='内嵌视频（add_movie），否则仅缩略图')
ap.add_argument('--thumb-w', type=float, default=0.94)
args = ap.parse_args()

SRC = os.path.abspath(args.ppt)
TMP = os.path.join(os.environ['TEMP'], 'ppe_work.pptx')
shutil.copyfile(SRC, TMP)  # OneDrive 脱水防护
spec = json.load(open(args.spec, encoding='utf-8'))
pr = Presentation(TMP)
slides = list(pr.slides)

def para_replace(slide, old, new):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            if old in p.text:
                runs = p.runs
                if runs:
                    runs[0].text = p.text.replace(old, new)
                    for r in runs[1:]:
                        r.text = ''
                return True
    return False

for page, cfg in spec.items():
    slide = slides[int(page) - 1]
    items = cfg['items']
    # 底条 y：页面内容最低点之下，至少 6.92in
    used = Emu(0)
    for sh in slide.shapes:
        try:
            bottom = sh.top + sh.height
            if bottom > used:
                used = bottom
        except Exception:
            pass
    strip_y = max(6.92, used / 914400.0 + 0.04)
    strip_h = 7.44 - strip_y
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(strip_y), Inches(13.333), Inches(strip_h))
    band.fill.solid(); band.fill.fore_color.rgb = DARK
    band.line.fill.background(); band.shadow.inherit = False
    # 路径说明（右）
    tb = slide.shapes.add_textbox(Inches(6.6), Inches(strip_y + 0.05), Inches(6.55), Inches(strip_h - 0.08))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]
    r1 = p.add_run(); r1.text = '检索过程录屏：'
    r1.font.size = Pt(9); r1.font.bold = True; r1.font.color.rgb = WHITE; r1.font.name = 'Microsoft YaHei'
    r2 = p.add_run(); r2.text = cfg.get('path_text', '')
    r2.font.size = Pt(8); r2.font.color.rgb = LIGHT; r2.font.name = 'Microsoft YaHei'
    p2 = tf.add_paragraph()
    r3 = p2.add_run(); r3.text = '逐字符真实键入检索式 → 点击检索 → 命中数出现（连续录像）'
    r3.font.size = Pt(7.5); r3.font.color.rgb = LIGHT; r3.font.name = 'Microsoft YaHei'
    # 缩略图 / 视频行
    x = 0.12
    tw = args.thumb_w
    for it in items:
        thumb = os.path.join(args.thumbs, it['key'] + '.jpg')
        th = strip_h - 0.10
        if args.embed_video:
            mp4 = os.path.join(cfg['videos_dir'], it['video'])
            poster = thumb if os.path.exists(thumb) else None
            mv = slide.shapes.add_movie(mp4, Inches(x), Inches(strip_y + 0.05), Inches(tw), Inches(th),
                                        poster_frame_image=poster, mime_type='video/mp4')
        else:
            im = Image.open(thumb)
            iw, ih = im.size
            sc = min(tw / iw, th / ih)
            iw2, ih2 = iw * sc, ih * sc
            pic = slide.shapes.add_picture(thumb, Inches(x + (tw - iw2) / 2), Inches(strip_y + 0.05),
                                           Inches(iw2), Inches(ih2))
            pic.line.color.rgb = WHITE; pic.line.width = Pt(0.75)
        lab = slide.shapes.add_textbox(Inches(x - 0.06), Inches(strip_y + strip_h - 0.20), Inches(tw + 0.12), Inches(0.19))
        pl = lab.text_frame.paragraphs[0]; pl.alignment = PP_ALIGN.CENTER
        rr = pl.add_run(); rr.text = it['label']
        rr.font.size = Pt(7); rr.font.bold = True; rr.font.color.rgb = WHITE; rr.font.name = 'Microsoft YaHei'
        x += tw + 0.06
    print(f'P{page}: 条 y={strip_y:.2f} 项 {len(items)}（{"视频内嵌" if args.embed_video else "缩略图"}）')

pr.save(TMP)
shutil.copyfile(TMP, SRC)
print('已写回:', SRC, f'({os.path.getsize(SRC):,} bytes)')
